"""Exercise extraction and retrieval against real embedded Qdrant."""

from dataclasses import replace
from io import BytesIO
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

import numpy as np
from PIL import Image
from qdrant_client import QdrantClient, models

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from content_cli import main as cli_main
from content_search.config import Settings
from content_search.diagnostics import diagnose
from content_search.extract import discover_files, extract_sections, split_text
from content_search.models import BGEEmbedder, Embedding, QwenAnalyzer
from content_search.pipeline import ContentPipeline
from content_search.store import ContentStore, match


class FixtureEmbedder:
    size = 1024
    hybrid = True
    identity = "fixture:keyword-space"

    def encode(self, texts):
        results = []
        for text in texts:
            dense = [0.0] * self.size
            sparse = {}
            for index, keyword in enumerate(("rocket", "garden", "timeout", "contract")):
                if keyword in text.casefold():
                    dense[index] = 1.0
                    sparse[index] = 1.0
            if not sparse:
                dense[4] = 1.0
                sparse[4] = 1.0
            results.append(Embedding(dense, sparse))
        return results

    def close(self):
        pass


class FixtureAnalyzer:
    def annotate(self, text="", image=None):
        if image is not None:
            return {"summary": "A timeout screenshot", "description": "Login page with a timeout error",
                    "ocr_text": "ERR_TIMEOUT 504", "tags": ["Screenshot", "Timeout"]}
        tags = [word for word in ("rocket", "garden", "contract") if word in text.casefold()]
        return {"summary": text[:200], "description": "", "ocr_text": "", "tags": tags}

    def answer(self, question, sources):
        return "Evidence [1]: " + sources[0]["text"]

    def close(self):
        pass


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)
        self.settings = Settings(qdrant_path=str(self.directory / "db"), collection="test",
                                 chunk_size=100, chunk_overlap=10, batch_size=2)
        self.embedder = FixtureEmbedder()
        self.analyzer = FixtureAnalyzer()
        self.store = ContentStore(self.settings, self.embedder)
        self.pipeline = ContentPipeline(self.settings, self.analyzer, self.embedder, self.store)

    def tearDown(self):
        self.pipeline.close()
        self.temp.cleanup()

    def write(self, name, text):
        path = self.directory / name
        path.write_text(text, encoding="utf-8")
        return path

    def test_hybrid_search_groups_files_and_applies_normalized_filters(self):
        self.pipeline.index_file(self.write("a.txt", "rocket launch notes\n" * 30))
        self.pipeline.index_file(self.write("b.md", "garden watering schedule"))
        self.pipeline.index_file(self.write("c.md", "rocket engine design"))
        results = self.pipeline.search("rocket", tags=[" ROCKET "], limit=10)
        self.assertEqual({item["name"] for item in results}, {"a.txt", "c.md"})
        self.assertEqual(len({item["file_id"] for item in results}), len(results))
        self.assertEqual(self.pipeline.search("rocket", tags=["garden", "rocket"]), [])
        self.assertEqual([item["name"] for item in self.pipeline.search("rocket", extension=".TXT", tags=["rocket"])], ["a.txt"])

    def test_unchanged_skips_models_and_update_removes_old_content(self):
        path = self.write("notes.txt", "rocket launch")
        self.pipeline.index_file(path)
        with patch.object(self.analyzer, "annotate", side_effect=AssertionError("Must skip models")):
            self.assertEqual(self.pipeline.index_file(path)["status"], "unchanged")
        path.write_text("garden notes", encoding="utf-8")
        self.pipeline.index_file(path)
        self.assertEqual(self.pipeline.search("rocket", tags=["rocket"]), [])
        self.assertEqual(len(self.pipeline.search("garden", tags=["garden"])), 1)
        points, _ = self.store.client.scroll("test", limit=100)
        self.assertEqual(len({p.payload["revision"] for p in points}), 1)

    def test_model_failure_preserves_searchable_previous_version(self):
        path = self.write("notes.txt", "rocket launch")
        self.pipeline.index_file(path)
        path.write_text("garden notes", encoding="utf-8")
        with patch.object(self.embedder, "encode", side_effect=RuntimeError("model unavailable")):
            result = self.pipeline.index(path)
        self.assertEqual(result[0]["status"], "failed")
        self.assertEqual(len(self.pipeline.search("rocket", tags=["rocket"])), 1)

    def test_failed_upload_stays_invisible_then_retry_cleans_staging(self):
        path = self.write("notes.txt", "rocket launch")
        self.pipeline.index_file(path)
        path.write_text("garden watering\n" * 30, encoding="utf-8")
        original = self.store.client.upsert
        calls = 0

        def fail_second(*args, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise RuntimeError("interrupted upload")
            return original(*args, **kwargs)

        with patch.object(self.store.client, "upsert", side_effect=fail_second):
            self.assertEqual(self.pipeline.index(path)[0]["status"], "failed")
        self.assertEqual(self.pipeline.search("garden", tags=["garden"]), [])
        self.assertEqual(len(self.pipeline.search("rocket", tags=["rocket"])), 1)
        self.pipeline.index_file(path)
        self.assertEqual(len(self.pipeline.search("garden", tags=["garden"])), 1)
        self.assertEqual(self.pipeline.search("rocket", tags=["rocket"]), [])
        self.assertEqual(self.store.client.count("test", count_filter=models.Filter(
            must=[match("ready", False)])).count, 0)

    def test_persistence_across_client_restart(self):
        self.pipeline.index_file(self.write("persist.txt", "rocket contract"))
        self.store.close()
        self.store = ContentStore(self.settings, self.embedder)
        self.pipeline.store = self.store
        self.assertEqual(len(self.pipeline.search("rocket", tags=["contract"])), 1)

    def test_image_ocr_and_description_are_independent_searchable_chunks(self):
        path = self.directory / "error.png"
        Image.new("RGB", (80, 40), "white").save(path)
        self.pipeline.index_file(path)
        points, _ = self.store.client.scroll("test", limit=100)
        payloads = {p.payload["kind"]: p.payload for p in points}
        self.assertEqual(payloads["ocr_text"]["text"], "ERR_TIMEOUT 504")
        self.assertEqual(payloads["ocr_text"]["page"], 1)
        self.assertIn("Login page", payloads["visual_description"]["text"])
        self.assertEqual(len(self.pipeline.search("timeout", tags=["screenshot"])), 1)

    def test_ask_returns_numbered_sources_and_skips_generation_without_hits(self):
        self.pipeline.index_file(self.write("notes.txt", "rocket launch"))
        result = self.pipeline.ask("rocket")
        self.assertIn("[1]", result["answer"])
        self.assertTrue(result["sources"][0]["path"].endswith("notes.txt"))
        with patch.object(self.analyzer, "answer", side_effect=AssertionError("Must not call")):
            self.assertEqual(self.pipeline.ask("garden", tags=["garden"])["sources"], [])

    def test_collection_rejects_embedding_identity_change(self):
        other = FixtureEmbedder()
        other.identity = "another-model"
        with self.assertRaisesRegex(ValueError, "different embedding"):
            ContentStore(self.settings, other, client=self.store.client)

    def test_original_text_is_kept_even_when_summary_omits_identifiers(self):
        path = self.write("notes.txt", "rocket part ZX-901 costs 123.45")
        with patch.object(self.analyzer, "annotate", return_value={
            "summary": "rocket notes", "description": "", "ocr_text": "", "tags": ["rocket"]}):
            self.pipeline.index_file(path)
        points, _ = self.store.client.scroll("test", limit=100)
        original = next(p.payload for p in points if p.payload["kind"] == "original_text")
        self.assertIn("ZX-901 costs 123.45", original["text"])

    def test_force_reindex_replaces_instead_of_accumulating_chunks(self):
        path = self.write("notes.txt", "rocket launch")
        self.pipeline.index_file(path)
        self.assertEqual(self.pipeline.index_file(path, force=True)["status"], "indexed")
        self.assertEqual(self.store.client.count("test").count, 2)

    def test_config_change_invalidates_unchanged_fingerprint(self):
        path = self.write("notes.txt", "rocket launch")
        self.pipeline.index_file(path)
        self.pipeline.settings = replace(self.settings, vl_model="changed-model")
        self.assertEqual(self.pipeline.index_file(path)["status"], "indexed")

    def test_dense_only_collection_search_respects_tags(self):
        embedder = FixtureEmbedder()
        embedder.hybrid = False
        settings = replace(self.settings, collection="dense_test")
        store = ContentStore(settings, embedder, client=self.store.client)
        pipeline = ContentPipeline(settings, self.analyzer, embedder, store)
        pipeline.index_file(self.write("dense.txt", "rocket launch"))
        self.assertEqual(len(pipeline.search("rocket", tags=["rocket"])), 1)
        self.assertEqual(pipeline.search("rocket", tags=["garden"]), [])


class ExtractionTests(unittest.TestCase):
    def test_chunking_keeps_every_character_without_infinite_overlap(self):
        text = "abcdef" * 200
        chunks = list(split_text(text, size=100, overlap=10))
        reconstructed = chunks[0] + "".join(chunk[10:] for chunk in chunks[1:])
        self.assertEqual(reconstructed, text)
        self.assertTrue(all(len(chunk) <= 100 for chunk in chunks))
        with self.assertRaises(ValueError):
            list(split_text(text, 100, 100))

    def test_pdf_preserves_page_numbers_and_reads_scan(self):
        import pymupdf
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "sample.pdf"
            document = pymupdf.open()
            document.new_page().insert_text((50, 50), "rocket contract ZX-901")
            image = BytesIO()
            Image.new("RGB", (100, 100), "white").save(image, format="PNG")
            page = document.new_page()
            page.insert_image(page.rect, stream=image.getvalue())
            document.save(path)
            document.close()
            sections = list(extract_sections(path))
            self.assertEqual(sections[0].page, 1)
            self.assertIn("ZX-901", sections[0].text)
            self.assertEqual(sections[1].page, 2)
            self.assertIsNotNone(sections[1].image)
            with self.assertRaisesRegex(ValueError, "enable PDF vision"):
                list(extract_sections(path, "never"))

    def test_docx_reads_paragraphs_tables_and_embedded_images(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            image = root / "picture.png"
            Image.new("RGB", (30, 30), "white").save(image)
            document = Document()
            document.add_paragraph("rocket contract")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "ZX-901"
            table.cell(0, 1).text = "123.45"
            document.add_picture(str(image))
            path = root / "document.docx"
            document.save(path)
            sections = list(extract_sections(path))
            self.assertIn("ZX-901\t123.45", sections[0].text)
            self.assertIsNotNone(sections[1].image)

    def test_discovery_skips_hidden_and_environment_directories(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / "visible.txt").write_text("hello")
            (root / ".env").write_text("private")
            (root / ".venv").mkdir()
            (root / ".venv" / "skip.txt").write_text("skip")
            self.assertEqual([p.name for p in discover_files(root)], ["visible.txt"])


class AdapterTests(unittest.TestCase):
    def test_doctor_checks_real_embedded_qdrant_and_reports_missing_models(self):
        with tempfile.TemporaryDirectory() as temp, patch("content_search.diagnostics.OpenAI") as client:
            client.return_value.__enter__.return_value.models.list.return_value.data = []
            checks = diagnose(Settings(qdrant_path=temp))
            self.assertFalse(checks[0]["ready"])
            self.assertTrue(checks[2]["ready"])

    def test_samantha_content_dispatch_preserves_quoted_arguments(self):
        from samantha import main
        with patch.object(sys, "argv", ["samantha.py", "content", "index", "path with spaces"]), \
                patch("content_cli.main", return_value=0) as content, patch("samantha.write_current_dir"):
            self.assertEqual(main(), 0)
            content.assert_called_once_with(["index", "path with spaces"])

    def test_vl_image_request_uses_encoded_image_and_json_output(self):
        with patch("content_search.models.OpenAI") as client:
            analyzer = QwenAnalyzer(Settings())
            response = Mock()
            response.choices = [Mock(finish_reason="stop", message=Mock(content=json.dumps({
                "summary": "timeout", "description": "login screen", "ocr_text": "504", "tags": [" ERROR "]})))]
            client.return_value.chat.completions.create.return_value = response
            image = BytesIO()
            Image.new("RGB", (20, 20), "white").save(image, format="PNG")
            data = analyzer.annotate(image=image.getvalue())
            self.assertEqual(data["tags"], ["error"])
            options = client.return_value.chat.completions.create.call_args.kwargs
            url = options["messages"][1]["content"][1]["image_url"]["url"]
            self.assertTrue(url.startswith("data:image/jpeg;base64,"))
            self.assertEqual(options["response_format"], {"type": "json_object"})
    def test_ollama_dense_adapter_disables_silent_truncation(self):
        settings = Settings(embedding_backend="ollama", embedding_model="bge-m3")
        embedder = BGEEmbedder(settings)
        try:
            with patch.object(embedder.client, "post") as request:
                request.return_value.json.return_value = {"embeddings": [[1.0] + [0.0] * 1023]}
                self.assertEqual(len(embedder.encode(["hello"])[0].dense), 1024)
                self.assertFalse(request.call_args.kwargs["json"]["truncate"])
                self.assertFalse(embedder.hybrid)
        finally:
            embedder.close()

    def test_flagembedding_adapter_converts_sparse_token_ids(self):
        embedder = BGEEmbedder(Settings())
        embedder.model = Mock()
        embedder.model.encode.return_value = {
            "dense_vecs": np.array([[1.0] + [0.0] * 1023]), "lexical_weights": [{"7": 0.5}]}
        try:
            self.assertEqual(embedder.encode(["hello"])[0].sparse, {7: 0.5})
            self.assertTrue(embedder.model.encode.call_args.kwargs["return_sparse"])
        finally:
            embedder.close()

    def test_invalid_vl_output_is_rejected_before_indexing(self):
        with patch("content_search.models.OpenAI") as client:
            analyzer = QwenAnalyzer(Settings())
            response = Mock()
            response.choices = [Mock(finish_reason="stop", message=Mock(content=json.dumps({
                "summary": "rocket", "description": "", "ocr_text": "", "tags": "invalid"})))]
            client.return_value.chat.completions.create.return_value = response
            with self.assertRaisesRegex(ValueError, "list of strings"):
                analyzer.annotate("rocket")

    def test_cli_propagates_partial_index_failure_as_nonzero_exit(self):
        with patch("content_cli.ContentPipeline") as pipeline:
            pipeline.return_value.index.return_value = [{"status": "failed", "path": "bad.pdf", "error": "bad file"}]
            with patch("builtins.print"):
                self.assertEqual(cli_main(["index", "bad.pdf", "--json"]), 1)
            pipeline.return_value.close.assert_called_once()


if __name__ == "__main__":
    unittest.main()
